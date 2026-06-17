import openpyxl as xl

# used to load the workbook
wb = xl.load_workbook('transactions-forpyautomation.xlsx')

sheet = wb['Sheet1']
# accessed the sheet 1 for the excel file

# cell = sheet['a1']
# we are definning the cell 
cell = sheet.cell(1,1) 
# better method to get the cell(row,column)

print(cell.value)

for row in range(2,sheet.max_row +1):
    cell = sheet.cell(row,3)
    corrected_price = cell.value *0.7
    correected_price_cell = sheet.cell(row,4)
    correected_price_cell.value = corrected_price
    #print(cell.value) # we got the price

wb.save('transactions-forpyautomation2(copy).xlsx')